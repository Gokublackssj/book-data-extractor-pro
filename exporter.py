"""Exportación de los libros limpios a CSV y Excel.

El exportador no sabe de dónde vienen los datos. Recibe una lista de
diccionarios con el contrato de ``config.COLUMNS`` y escribe archivos.
"""

from __future__ import annotations

import csv
import logging
import re
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import COLUMNS, COLUMN_ORDER, Config

logger = logging.getLogger(__name__)


class ExportError(RuntimeError):
    """Error al escribir los archivos de salida."""


@dataclass(frozen=True)
class ExportResult:
    """Rutas generadas y número de filas exportadas."""

    csv_path: Path | None
    xlsx_path: Path | None
    filas: int


# Caracteres que Excel puede interpretar como inicio de fórmula en un CSV.
_PREFIJOS_PELIGROSOS = ("=", "+", "-", "@", "\t", "\r")

_ANCHOS_COLUMNA = {
    "titulo": 45,
    "precio": 12,
    "moneda": 10,
    "disponibilidad": 16,
    "url": 60,
    "advertencias": 40,
}


def _sanitizar_prefijo(prefijo: str) -> str:
    """Evita que un prefijo con ``../`` escriba fuera de output/."""

    limpio = re.sub(r"[^A-Za-z0-9._-]+", "_", prefijo).strip("._")
    return limpio or "libros"


def _ruta_base(config: Config) -> Path:
    """Genera una ruta con fecha y hora que no sobrescriba archivos."""

    try:
        config.output_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ExportError(
            f"No se pudo crear la carpeta de salida {config.output_dir}: {exc}"
        ) from exc

    prefijo = _sanitizar_prefijo(config.prefix)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = config.output_dir / f"{prefijo}_{sello}"

    contador = 0
    while (
        base.parent / f"{base.name}.csv"
    ).exists() or (
        base.parent / f"{base.name}.xlsx"
    ).exists():
        contador += 1
        base = config.output_dir / f"{prefijo}_{sello}_{contador:02d}"

    return base


def _valor_celda(libro: dict, campo: str):
    """Convierte un campo interno al valor que se escribe en el archivo."""

    if campo == "precio":
        precio = libro.get("precio")
        return float(precio) if precio is not None else None

    if campo == "advertencias":
        return "; ".join(libro.get("advertencias") or [])

    return libro.get(campo) or ""


def _celda_csv(valor):
    """Mitiga la inyección de fórmulas al abrir el CSV con Excel."""

    if isinstance(valor, str) and valor.startswith(_PREFIJOS_PELIGROSOS):
        return "'" + valor
    return "" if valor is None else valor


def _escribir_csv(libros: Sequence[dict], ruta: Path, config: Config) -> None:
    try:
        with ruta.open("w", encoding="utf-8-sig", newline="") as archivo:
            escritor = csv.writer(archivo, delimiter=config.csv_delimiter)
            escritor.writerow([COLUMNS[campo] for campo in COLUMN_ORDER])
            for libro in libros:
                escritor.writerow(
                    [
                        _celda_csv(_valor_celda(libro, campo))
                        for campo in COLUMN_ORDER
                    ]
                )
    except PermissionError as exc:
        raise ExportError(
            f"No se pudo escribir {ruta}. "
            "Si el archivo está abierto en Excel, ciérralo e intenta de nuevo."
        ) from exc
    except OSError as exc:
        raise ExportError(f"No se pudo escribir el CSV {ruta}: {exc}") from exc


def _escribir_xlsx(libros: Sequence[dict], ruta: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Libros"

    ws.append([COLUMNS[campo] for campo in COLUMN_ORDER])

    for libro in libros:
        fila = [_valor_celda(libro, campo) for campo in COLUMN_ORDER]
        ws.append(fila)

        # openpyxl interpreta como fórmula cualquier cadena que empiece por "=".
        # Forzamos el tipo texto para que un título malicioso no se ejecute.
        for indice, valor in enumerate(fila, start=1):
            if isinstance(valor, str) and valor.startswith("="):
                ws.cell(row=ws.max_row, column=indice).data_type = "s"

    # Estilo del encabezado.
    fuente_encabezado = Font(bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill("solid", fgColor="1F4E78")
    for celda in ws[1]:
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center")

    # Anchos, filtro y panel congelado para que sea cómodo de leer.
    for indice, campo in enumerate(COLUMN_ORDER, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = _ANCHOS_COLUMNA.get(
            campo, 20
        )

    columna_precio = COLUMN_ORDER.index("precio") + 1
    for fila in range(2, ws.max_row + 1):
        ws.cell(row=fila, column=columna_precio).number_format = "#,##0.00"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    try:
        wb.save(ruta)
    except PermissionError as exc:
        raise ExportError(
            f"No se pudo escribir {ruta}. "
            "Si el archivo está abierto en Excel, ciérralo e intenta de nuevo."
        ) from exc
    except OSError as exc:
        raise ExportError(f"No se pudo escribir el Excel {ruta}: {exc}") from exc


def export_books(libros: Sequence[dict], config: Config) -> ExportResult:
    """Exporta los libros según ``config.export_format``."""

    base = _ruta_base(config)
    csv_path: Path | None = None
    xlsx_path: Path | None = None

    if config.export_format in {"both", "csv"}:
        csv_path = base.parent / f"{base.name}.csv"
        _escribir_csv(libros, csv_path, config)
        logger.info("CSV generado: %s", csv_path)

    if config.export_format in {"both", "xlsx"}:
        xlsx_path = base.parent / f"{base.name}.xlsx"
        _escribir_xlsx(libros, xlsx_path)
        logger.info("Excel generado: %s", xlsx_path)

    return ExportResult(csv_path=csv_path, xlsx_path=xlsx_path, filas=len(libros))
