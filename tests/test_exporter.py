"""Pruebas de exportación a CSV y Excel."""

import csv
from decimal import Decimal

from openpyxl import load_workbook

from config import Config
from exporter import export_books


def _libros_de_ejemplo():
    return [
        {
            "titulo": "A Light in the Attic",
            "precio": Decimal("51.77"),
            "moneda": "£",
            "disponibilidad": "disponible",
            "url": "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html",
            "advertencias": [],
        },
        {
            "titulo": "=1+1",
            "precio": None,
            "moneda": "",
            "disponibilidad": "desconocida",
            "url": "",
            "advertencias": ["precio ausente o inválido"],
        },
    ]


def _leer_csv(ruta, delimitador=","):
    with ruta.open(encoding="utf-8-sig", newline="") as archivo:
        return list(csv.reader(archivo, delimiter=delimitador))


def test_exporta_csv_y_xlsx(tmp_path):
    config = Config(output_dir=tmp_path, prefix="prueba", export_format="both")

    resultado = export_books(_libros_de_ejemplo(), config)

    assert resultado.filas == 2
    assert resultado.csv_path is not None
    assert resultado.csv_path.exists()
    assert resultado.xlsx_path is not None
    assert resultado.xlsx_path.exists()

    filas = _leer_csv(resultado.csv_path)
    assert filas[0] == [
        "Título",
        "Precio",
        "Moneda",
        "Disponibilidad",
        "URL",
        "Advertencias",
    ]
    assert filas[1][0] == "A Light in the Attic"
    assert filas[1][1] == "51.77"
    # La celda peligrosa se neutraliza en CSV.
    assert filas[2][0] == "'=1+1"

    wb = load_workbook(resultado.xlsx_path)
    ws = wb.active
    assert ws["A1"].value == "Título"
    assert ws["A2"].value == "A Light in the Attic"
    assert ws["B2"].number_format == "#,##0.00"
    assert ws["A3"].value == "=1+1"
    assert ws["A3"].data_type == "s"  # Texto, no fórmula.


def test_no_sobrescribe_archivos_anteriores(tmp_path):
    config = Config(output_dir=tmp_path, prefix="unico", export_format="both")

    primero = export_books(_libros_de_ejemplo(), config)
    segundo = export_books(_libros_de_ejemplo(), config)

    assert primero.csv_path != segundo.csv_path
    assert primero.xlsx_path != segundo.xlsx_path
    assert primero.csv_path.exists()
    assert segundo.csv_path.exists()


def test_solo_csv(tmp_path):
    config = Config(output_dir=tmp_path, prefix="solo_csv", export_format="csv")

    resultado = export_books(_libros_de_ejemplo(), config)

    assert resultado.csv_path is not None
    assert resultado.csv_path.exists()
    assert resultado.xlsx_path is None


def test_solo_xlsx(tmp_path):
    config = Config(output_dir=tmp_path, prefix="solo_xlsx", export_format="xlsx")

    resultado = export_books(_libros_de_ejemplo(), config)

    assert resultado.csv_path is None
    assert resultado.xlsx_path is not None
    assert resultado.xlsx_path.exists()


def test_lista_vacia_escribe_solo_encabezados(tmp_path):
    config = Config(output_dir=tmp_path, prefix="vacio", export_format="both")

    resultado = export_books([], config)

    filas = _leer_csv(resultado.csv_path)
    assert len(filas) == 1
    assert filas[0][0] == "Título"

    wb = load_workbook(resultado.xlsx_path)
    assert wb.active.max_row == 1


def test_prefijo_peligroso_no_escribe_fuera_de_output(tmp_path):
    config = Config(
        output_dir=tmp_path,
        prefix="../../fuera",
        export_format="csv",
    )

    resultado = export_books(_libros_de_ejemplo(), config)

    assert resultado.csv_path.parent == tmp_path
    assert resultado.csv_path.exists()


def test_delimitador_personalizado(tmp_path):
    config = Config(
        output_dir=tmp_path,
        prefix="punto_y_coma",
        export_format="csv",
        csv_delimiter=";",
    )

    resultado = export_books(_libros_de_ejemplo(), config)
    filas = _leer_csv(resultado.csv_path, delimitador=";")

    assert len(filas[0]) == 6
    assert filas[0][0] == "Título"
